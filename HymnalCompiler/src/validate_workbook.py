"""
validate_workbook.py
 
Stage 1 of the Hymnal Edition 2 Typst pilot.
 
Reads the workbook (tabs "Hymns Metadata", "Lyrics Section", "Contributors")
and validates it against the established data model. Prints a clear
ERRORS / WARNINGS / INFO report. Does NOT modify the workbook and does
NOT produce a build (.typ/.pdf) — that's Stage 2.
 
SCHEMA NOTE (updated): Author, Translator, Composer 1-3 have been removed
from Hymns Metadata and normalized into the "Contributors" sheet (one row
per credited person/source per hymn). Hymns Metadata gained a "Main Tune"
column (blank = Tune 1 is primary, or 1/2/3 to say otherwise).
 
Usage:
    python3 validate_workbook.py "Hymns We Sing 2nd Edition Master - Pilot.xlsx"
"""
 
import sys
import re
from collections import defaultdict
import openpyxl
from workbook_schema import read_title_page
 
VALID_TYPES = {"stanza", "chorus", "intro", "outro", "bridge", "final chorus"}
VALID_STATUSES = {"active", "inactive", "removed"}
 
HYMNS_SHEET = "Hymns Metadata"
SECTIONS_SHEET = "Lyrics Section"
CONTRIBUTORS_SHEET = "Contributors"
CATEGORIES_SHEET = "Categories"
TITLE_PAGE_SHEET = "Title Page"
 
HYMNS_COLUMNS = [
    "ID", "Title", "Category", "Tune 1", "Tune 2", "Tune 3", "Main Tune", "Status",
]
SECTIONS_COLUMNS = ["Hymn ID", "Sequence", "Type", "Label", "Text"]
CONTRIBUTORS_COLUMNS = ["Hymn ID", "Sequence", "Role", "Person Name", "Note", "Tune Slot"]
CATEGORIES_COLUMNS = ["Category", "Sequence"]

# Check these fields before domain validation calls string methods. Numeric
# sequences, tune slots, Main Tune and stanza labels remain supported.
TEXT_COLUMNS = {
    HYMNS_SHEET: {"ID", "Title", "Category", "Tune 1", "Tune 2", "Tune 3", "Status"},
    SECTIONS_SHEET: {"Hymn ID", "Type", "Text"},
    CONTRIBUTORS_SHEET: {"Hymn ID", "Role", "Person Name", "Note"},
    CATEGORIES_SHEET: {"Category"},
}
 
VALID_ROLES = {"author", "word source", "translator", "translation source", "composer", "tune source"}
# Which of the three Sequence-numbering groups each Role belongs to (established
# 2026-09-24): (author, word source) / (translator, translation source) /
# (composer, tune source). Sequence uniqueness is checked within a group, not
# across the whole hymn — a "Sequence 1" author and a "Sequence 1" translator
# are not a collision.
ROLE_GROUP = {
    "author": "words", "word source": "words",
    "translator": "translation", "translation source": "translation",
    "composer": "tune", "tune source": "tune",
# For the tune group, Sequence uniqueness is also checked per Tune Slot.
}
TUNE_SLOT_ROLES = {"composer", "tune source"}
VALID_TUNE_SLOTS = (1, 2, 3, "1", "2", "3")
 
 
class Report:
    def __init__(self):
        self.errors = []
        self.warnings = []
        self.info = []
 
    def error(self, msg):
        self.errors.append(msg)
 
    def warning(self, msg):
        self.warnings.append(msg)
 
    def note(self, msg):
        self.info.append(msg)
 
    def print_all(self):
        def section(title, items, marker):
            print("=" * 70)
            print(f"{title} ({len(items)})")
            print("=" * 70)
            if not items:
                print("  none")
            for m in items:
                print(f"  {marker} {m}")
            print()
 
        section("ERRORS — must fix before build", self.errors, "\u2717")
        section("WARNINGS — review before build", self.warnings, "!")
        section("INFO — for awareness only", self.info, "\u00b7")
        print(f"Summary: {len(self.errors)} error(s), {len(self.warnings)} warning(s)")
 
 
def load_sheet_rows(ws, expected_columns, sheet_label, report):
    """Read a worksheet into a list of dicts keyed by header name."""
    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    header_row = [h.strip() if isinstance(h, str) else h for h in header_row]
 
    missing = [c for c in expected_columns if c not in header_row]
    extra = [h for h in header_row if h and h not in expected_columns]
    if missing:
        report.error(f"{sheet_label}: missing expected column(s): {missing}")
    if extra:
        report.note(f"{sheet_label}: unexpected extra column(s) present (ignored): {extra}")
 
    col_index = {name: idx for idx, name in enumerate(header_row) if name in expected_columns}
 
    rows = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2), start=2):
        values = [c.value for c in row]
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in values):
            continue  # skip fully blank rows
        record = {}
        for col_name in expected_columns:
            idx = col_index.get(col_name)
            val = values[idx] if idx is not None and idx < len(values) else None
            if isinstance(val, str):
                val = val.strip()
            if col_name in TEXT_COLUMNS.get(sheet_label, ()) and val is not None and not isinstance(val, str):
                report.error(
                    f"{sheet_label} row {row_num}, {col_name}: "
                    f"expected text, got {type(val).__name__}"
                )
            record[col_name] = val
        record["_row"] = row_num
        rows.append(record)
    return rows
 
 
def id_from_title(title):
    """First letter of every word (incl. a/an/the/of/and/to), punctuation
    ignored, uppercase — the established ID rule. Straight and curly
    apostrophes stay inside a word, but a leading apostrophe (e.g. 'Tis)
    is not counted as the word's initial."""
    words = re.findall(r"[A-Za-z'\u2018\u2019]+", title)
    initials = []
    for w in words:
        w = w.lstrip("'\u2018\u2019")
        if w:
            initials.append(w[0].upper())
    return "".join(initials)
 
 
def validate_hymns(hymns, report):
    ids_seen = {}
    titles_seen = defaultdict(list)
    valid_ids = set()
 
    for h in hymns:
        row, hid, title = h["_row"], h["ID"], h["Title"]
 
        if not hid:
            report.error(f"Hymns Metadata row {row}: missing ID")
            continue
        if not title:
            report.error(f"Hymns Metadata row {row} (ID {hid}): missing Title")
 
        if hid in ids_seen:
            report.error(f"Hymns Metadata: duplicate ID '{hid}' at rows {ids_seen[hid]} and {row}")
        else:
            ids_seen[hid] = row
        valid_ids.add(hid)
 
        if not h["Category"]:
            report.error(f"Hymns Metadata row {row} (ID {hid}): missing Category")
 
        status = (h["Status"] or "").lower()
        if not status:
            report.error(f"Hymns Metadata row {row} (ID {hid}): missing Status")
        elif status not in VALID_STATUSES:
            report.error(
                f"Hymns Metadata row {row} (ID {hid}): invalid Status '{h['Status']}' "
                f"(expected one of {sorted(VALID_STATUSES)})"
            )
 
        t1, t2, t3 = h.get("Tune 1"), h.get("Tune 2"), h.get("Tune 3")
 
        if not t1 and (t2 or t3):
            report.error(
                f"Hymns Metadata row {row} (ID {hid}): Tune 1 is blank but Tune 2/3 filled "
                f"— tune order (1 = primary unless Main Tune says otherwise) must be preserved"
            )
        if not t2 and t3:
            report.error(
                f"Hymns Metadata row {row} (ID {hid}): Tune 2 is blank but Tune 3 filled "
                f"— tune order must be preserved"
            )
 
        tune_names = [t for t in (t1, t2, t3) if t]
        if len(tune_names) != len(set(tune_names)):
            report.warning(f"Hymns Metadata row {row} (ID {hid}): repeated tune name across Tune 1-3: {tune_names}")
 
        main_tune = h.get("Main Tune")
        if main_tune not in (None, "", *VALID_TUNE_SLOTS):
            report.error(
                f"Hymns Metadata row {row} (ID {hid}): invalid Main Tune '{main_tune}' "
                f"(expected blank, 1, 2, or 3)"
            )
        elif main_tune not in (None, ""):
            slot_num = int(main_tune)
            slot_tune = {1: t1, 2: t2, 3: t3}[slot_num]
            if not slot_tune:
                report.error(
                    f"Hymns Metadata row {row} (ID {hid}): Main Tune points to Tune {slot_num}, "
                    f"but Tune {slot_num} is blank"
                )
 
        if title:
            expected = id_from_title(title)
            base_id = re.sub(r"-\d+$", "", hid) if hid else ""
            if base_id and expected and base_id != expected:
                report.note(
                    f"Hymns Metadata row {row}: ID '{hid}' doesn't match title-derived ID "
                    f"'{expected}' for title '{title}' — OK if the title changed after the ID "
                    f"was assigned (IDs are permanent), otherwise verify"
                )
 
        if title:
            titles_seen[title.lower()].append((hid, row))
 
    for title_lc, entries in titles_seen.items():
        if len(entries) > 1:
            report.warning(f"Duplicate title (case-insensitive) '{title_lc}' used by: {entries}")
 
    return valid_ids
 
 
def validate_contributors(contributors, valid_ids, hymns_by_id, report):
    """Validates the Contributors sheet. Returns {hymn_id: [row, ...]}."""
    seq_seen = defaultdict(dict)  # (hid, group, tune_slot) -> {seq: row_num}
    by_hymn = defaultdict(list)
 
    for c in contributors:
        row, hid = c["_row"], c["Hymn ID"]
        if not hid:
            report.error(f"Contributors row {row}: missing Hymn ID")
            continue
        if hid not in valid_ids:
            report.error(f"Contributors row {row}: Hymn ID '{hid}' has no matching row in Hymns Metadata")
            continue
 
        seq = c["Sequence"]
        if seq is None or seq == "":
            report.error(f"Contributors row {row} (Hymn {hid}): missing Sequence")
        elif not isinstance(seq, (int, float)):
            report.error(f"Contributors row {row} (Hymn {hid}): Sequence '{seq}' is not numeric")
 
        role_raw = c["Role"]
        role = (role_raw or "").strip().lower()
        if not role:
            report.error(f"Contributors row {row} (Hymn {hid}): missing Role")
            continue
        if role not in VALID_ROLES:
            report.error(
                f"Contributors row {row} (Hymn {hid}): invalid Role '{role_raw}' "
                f"(expected one of {sorted(VALID_ROLES)})"
            )
            continue
 
        if not c["Person Name"]:
            report.error(f"Contributors row {row} (Hymn {hid}, Role {role}): missing Person Name")
 
        slot = c["Tune Slot"]
        slot_num = None
        if role in TUNE_SLOT_ROLES:
            if slot is None or slot == "":
                report.error(f"Contributors row {row} (Hymn {hid}): Role '{role}' requires a Tune Slot")
            elif slot not in VALID_TUNE_SLOTS:
                report.error(
                    f"Contributors row {row} (Hymn {hid}): invalid Tune Slot '{slot}' "
                    f"(expected 1, 2, or 3)"
                )
            else:
                slot_num = int(slot)
                h = hymns_by_id.get(hid)
                slot_tune = h.get(f"Tune {slot_num}") if h else None
                if h is not None and not slot_tune:
                    report.error(
                        f"Contributors row {row} (Hymn {hid}): Tune Slot {slot_num} given, "
                        f"but Tune {slot_num} is blank in Hymns Metadata"
                    )
        elif slot not in (None, ""):
            report.error(
                f"Contributors row {row} (Hymn {hid}): Role '{role}' should not have a Tune Slot "
                f"(only composer/tune source use it)"
            )
 
        by_hymn[hid].append(c)
 
        if isinstance(seq, (int, float)):
            key = (hid, ROLE_GROUP[role], slot_num)
            existing = seq_seen[key].get(seq)
            if existing is not None:
                report.error(
                    f"Hymn {hid}: duplicate Sequence {seq} within the same group "
                    f"('{ROLE_GROUP[role]}'" + (f", Tune Slot {slot_num}" if slot_num else "") +
                    f") — rows {existing} and {row}"
                )
            else:
                seq_seen[key][seq] = row
 
    return by_hymn
 
def validate_categories(categories, hymns, report):
    """Validates the Categories sheet. Returns {category: sequence}."""
    order = {}
    seq_seen = {}
    for c in categories:
        row, cat, seq = c["_row"], c["Category"], c["Sequence"]
        if not cat:
            report.error(f"Categories row {row}: missing Category")
            continue
        if seq is None or seq == "":
            report.error(f"Categories row {row} (Category '{cat}'): missing Sequence")
            continue
        if not isinstance(seq, (int, float)):
            report.error(f"Categories row {row} (Category '{cat}'): Sequence '{seq}' is not numeric")
            continue
        if cat in order:
            report.error(f"Categories: duplicate Category '{cat}' (row {row})")
        else:
            order[cat] = seq
        if seq in seq_seen:
            report.warning(f"Categories: Sequence {seq} used by both '{seq_seen[seq]}' and '{cat}' — tie order undefined")
        else:
            seq_seen[seq] = cat
 
    used = {(h["Category"] or "").strip() for h in hymns
            if (h["Status"] or "").lower() == "active" and h["Category"]}
    for cat in sorted(used - set(order)):
        report.error(f"Category '{cat}' is used by active hymn(s) but has no row in Categories sheet")
    for cat in sorted(set(order) - used):
        report.note(f"Categories: '{cat}' has a Sequence entry but no active hymn uses it")
 
    return order
 
def validate_hymn_contributor_coverage(hymns, by_hymn_contrib, report):
    """Parity with the old per-field checks: active hymns should have an
    Author/Word Source, and every populated tune slot should have a
    Composer/Tune Source — reported as warnings, not auto-filled."""
    for h in hymns:
        hid = h["ID"]
        if (h["Status"] or "").lower() != "active":
            continue
 
        rows = by_hymn_contrib.get(hid, [])
        roles_present = {(r["Role"] or "").strip().lower() for r in rows}
 
        if not (roles_present & {"author", "word source"}):
            report.warning(
                f"Hymn {hid}: no Author/Word Source in Contributors — flagged for review, "
                f"not auto-filled (spec: do not invent uncertain attribution)"
            )
 
        for n, tune in ((1, h.get("Tune 1")), (2, h.get("Tune 2")), (3, h.get("Tune 3"))):
            if not tune:
                continue
            has_composer = any(
                (r["Role"] or "").strip().lower() in TUNE_SLOT_ROLES and r["Tune Slot"] in (n, str(n))
                for r in rows
            )
            if not has_composer:
                report.warning(f"Hymn {hid}: Tune {n} '{tune}' has no Composer/Tune Source in Contributors")
 
 
def validate_sections(sections, valid_ids, report):
    by_hymn = defaultdict(list)
    for s in sections:
        row, hid = s["_row"], s["Hymn ID"]
        if not hid:
            report.error(f"Lyrics Section row {row}: missing Hymn ID")
            continue
        if hid not in valid_ids:
            report.error(f"Lyrics Section row {row}: Hymn ID '{hid}' has no matching row in Hymns Metadata")
            continue
 
        seq = s["Sequence"]
        if seq is None or seq == "":
            report.error(f"Lyrics Section row {row} (Hymn {hid}): missing Sequence")
        elif not isinstance(seq, (int, float)):
            report.error(f"Lyrics Section row {row} (Hymn {hid}): Sequence '{seq}' is not numeric")
 
        type_val = (s["Type"] or "").strip().lower()
        if not type_val:
            report.error(f"Lyrics Section row {row} (Hymn {hid}): missing Type")
        elif type_val not in VALID_TYPES:
            report.error(
                f"Lyrics Section row {row} (Hymn {hid}): invalid Type '{s['Type']}' "
                f"(expected Stanza, Chorus, Intro, Outro, Bridge, or Final Chorus)"
            )
 
        if not s["Text"]:
            report.error(f"Lyrics Section row {row} (Hymn {hid}, Type {s['Type']}): missing Text")
 
        by_hymn[hid].append(s)
 
    for hid, secs in by_hymn.items():
        seqs = [s["Sequence"] for s in secs if isinstance(s["Sequence"], (int, float))]
        if len(seqs) != len(set(seqs)):
            report.error(f"Hymn {hid}: duplicate Sequence values in Lyrics Section: {seqs}")
 
    for hid in valid_ids:
        if hid not in by_hymn:
            report.warning(f"Hymn {hid}: no rows in Lyrics Section at all (no lyric content)")
 
    return by_hymn
 
 
def validate_active_inactive(hymns, by_hymn, report):
    for h in hymns:
        hid = h["ID"]
        status = (h["Status"] or "").lower()
        if status == "active" and hid not in by_hymn:
            report.error(f"Hymn {hid}: Status=active but has no Lyrics Section content — cannot be published as-is")

# Straight ASCII quotes are flagged in published editorial text. Not
# auto-converted: the editor fixes the workbook by hand.
#   \u2019 = ’   \u2018 = ‘   \u201c = “   \u201d = ”
STRAIGHT_QUOTES = (
    ("'", "apostrophe/single quote",
     "use \u2019 (U+2019) for an apostrophe or closing quote, \u2018 (U+2018) for an opening quote"),
    ('"', "double quote",
     "use \u201c (U+201C) to open and \u201d (U+201D) to close"),
)
EXCERPT_RADIUS = 20


def _excerpt(text, pos):
    """Short single-line snippet of text around position pos."""
    start = max(0, pos - EXCERPT_RADIUS)
    end = min(len(text), pos + EXCERPT_RADIUS + 1)
    snippet = text[start:end].replace("\r", " ").replace("\n", " / ")
    return ("\u2026" if start > 0 else "") + snippet + ("\u2026" if end < len(text) else "")


def check_typography(value, where, report):
    """Warns once per straight-quote type found in one cell's text."""
    if not isinstance(value, str):
        return
    for char, name, advice in STRAIGHT_QUOTES:
        count = value.count(char)
        if count:
            report.warning(
                f"{where}: {count} straight {name}(s) ({char}) - {advice}. "
                f"First at: \u00ab{_excerpt(value, value.index(char))}\u00bb"
            )


def validate_typography(hymns, sections, contributors, title_page_ws, report):
    """Warns about straight ' and \" in published editorial text of ACTIVE
    hymns. Technical fields (IDs, Sequence, Type, Role, Status, Tune Slot,
    Main Tune) are deliberately not scanned."""
    titles = {h["ID"]: h["Title"] for h in hymns
              if h["ID"] and (h["Status"] or "").lower() == "active"}

    for h in hymns:
        if h["ID"] not in titles:
            continue
        where = f"{HYMNS_SHEET} row {h['_row']} (ID {h['ID']}, '{h['Title']}')"
        for field in ("Title", "Category", "Tune 1", "Tune 2", "Tune 3"):
            check_typography(h.get(field), f"{where}, {field}", report)

    for s in sections:
        hid = s["Hymn ID"]
        if hid not in titles:
            continue
        where = f"{SECTIONS_SHEET} row {s['_row']} (Hymn {hid}, '{titles[hid]}', {s['Type']})"
        for field in ("Label", "Text"):
            check_typography(s.get(field), f"{where}, {field}", report)

    for c in contributors:
        hid = c["Hymn ID"]
        if hid not in titles:
            continue
        where = f"{CONTRIBUTORS_SHEET} row {c['_row']} (Hymn {hid}, '{titles[hid]}', {c['Role']})"
        for field in ("Person Name", "Note"):
            check_typography(c.get(field), f"{where}, {field}", report)

    for row_num, (key, value) in enumerate(
        title_page_ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2
    ):
        if key:
            check_typography(value, f"{TITLE_PAGE_SHEET} row {row_num} (key '{key}')", report)
 
 
def main():
    if len(sys.argv) != 2:
        print("Usage: python3 validate_workbook.py <path-to-HYMNAL_EDITION_2.xlsx>")
        sys.exit(1)
 
    path = sys.argv[1]
    report = Report()
    wb = openpyxl.load_workbook(path, data_only=True)
 
    for sheet in (HYMNS_SHEET, SECTIONS_SHEET, CONTRIBUTORS_SHEET, CATEGORIES_SHEET, TITLE_PAGE_SHEET):
        if sheet not in wb.sheetnames:
            report.error(f"Workbook is missing sheet '{sheet}'. Found: {wb.sheetnames}")
    if report.errors:
        report.print_all()
        sys.exit(1)
 
    hymns = load_sheet_rows(wb[HYMNS_SHEET], HYMNS_COLUMNS, HYMNS_SHEET, report)
    sections = load_sheet_rows(wb[SECTIONS_SHEET], SECTIONS_COLUMNS, SECTIONS_SHEET, report)
    contributors = load_sheet_rows(wb[CONTRIBUTORS_SHEET], CONTRIBUTORS_COLUMNS, CONTRIBUTORS_SHEET, report)
    category_rows = load_sheet_rows(wb[CATEGORIES_SHEET], CATEGORIES_COLUMNS, CATEGORIES_SHEET, report)
    report.note(
        f"Loaded {len(hymns)} hymn row(s), {len(sections)} section row(s), "
        f"{len(contributors)} contributor row(s)"
    )

    _, title_errors = read_title_page(wb[TITLE_PAGE_SHEET])
    for error in title_errors:
        report.error(error)

    # Schema errors must stop here: the checks below assume text fields are
    # strings and would otherwise hide these useful messages in a traceback.
    if report.errors:
        report.print_all()
        sys.exit(1)
 
    valid_ids = validate_hymns(hymns, report)
    hymns_by_id = {h["ID"]: h for h in hymns if h["ID"]}
 
    by_hymn_contrib = validate_contributors(contributors, valid_ids, hymns_by_id, report)
    validate_hymn_contributor_coverage(hymns, by_hymn_contrib, report)
 
    by_hymn_sections = validate_sections(sections, valid_ids, report)
    validate_active_inactive(hymns, by_hymn_sections, report)
    validate_typography(hymns, sections, contributors, wb[TITLE_PAGE_SHEET], report)
 
    categories = defaultdict(int)
    for h in hymns:
        categories[h["Category"] or "(blank)"] += 1
    report.note("Categories in use: " + ", ".join(f"{c} ({n})" for c, n in sorted(categories.items())))
 
    validate_categories(category_rows, hymns, report)
 
    statuses = defaultdict(int)
    for h in hymns:
        s = (h["Status"] or "").strip().lower()
        statuses[s or "(blank)"] += 1
    report.note("Status breakdown: " + ", ".join(f"{s} ({n})" for s, n in sorted(statuses.items())))
 
    report.print_all()
    sys.exit(1 if report.errors else 0)
 
 
if __name__ == "__main__":
    main()
