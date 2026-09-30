"""
transform.py

Stage 2 of the Hymnal Edition 2 Typst pilot.

Loads the workbook, filters to ACTIVE hymns, assigns printed hymn numbers,
generates attribution text from the Contributors sheet, orders sections,
and builds the hymn data that main.typ/template.typ consume. Prints a
readable summary so results can be checked BEFORE any Typst source is
generated (that's Stage 3). Assumes validate_workbook.py has already been
run and its warnings reviewed.

SCHEMA NOTE (updated): Author, Translator, Composer 1-3 no longer exist as
flat columns on Hymns Metadata — they're rows in "Contributors" (Hymn ID,
Sequence, Role, Person Name, Note, Tune Slot). Hymns Metadata gained
"Main Tune" (blank = Tune 1 is primary, or 1/2/3 to override).

build_indexes() produces the four publication indexes (see its own docstring).

DECIDED for this pilot (confirmed by editor):

  - Printed Hymn No. sort order: strict case-insensitive alphabetical by
    Title, including leading articles (A/An/The are NOT ignored). Ties
    broken by ID.

  - Attribution name-joining: Oxford-comma list ("A, B, and C"; "A and B"
    for exactly two). A Contributors row's Note (stanza attribution or a
    language note) is appended in parentheses after that one name only.

  - By/from wording (2026-09-24): within each category, person-role credits
    (author / translator / composer) are joined under "by"; source-role
    credits (word source / translation source / tune source) are joined
    under "from"; the two halves are combined with " & " only when both are
    present, e.g. "Words by A and B & from C". Applies identically to Words,
    Translation, and Tune/Alt. tune (the tune credit is appended straight
    after the quoted tune name, e.g. 'Tune: "NETTLETON" by X & from Y').

  - "Words and Tune by/from X" collapse: applies when the Words credits
    and the Main Tune credits are the same set of names (one or several),
    all of the same kind (all person roles or all source roles), and none
    carries a Note. A Note signals a distinction (e.g. a stanza-only
    credit) that collapsing would erase. The connector ("by" or "from")
    follows the shared kind.

Usage:
    python transform.py <Hymns We Sing 2nd Edition Master Datasheet (LIVE)>
    """

import sys
from collections import defaultdict
import openpyxl
import re

HYMNS_SHEET = "Hymns Metadata"
SECTIONS_SHEET = "Lyrics Section"
CONTRIBUTORS_SHEET = "Contributors"
CATEGORIES_SHEET = "Categories"
TITLE_PAGE_SHEET = "Title Page"

HYMNS_COLUMNS = [
    "ID", "Title", "Category", "Tune 1", "Tune 2", "Tune 3", "Main Tune", "Status",
]
SECTIONS_COLUMNS = ["Hymn ID", "Sequence", "Type", "Label", "Text"]
CONTRIBUTORS_COLUMNS = ["Hymn ID", "Sequence", "Role", "Person Name", "Note", "Tune Slot"]

WORDS_ROLES = ("author", "word source")
TRANSLATION_ROLES = ("translator", "translation source")
TUNE_ROLES = ("composer", "tune source")

CATEGORIES_COLUMNS = ["Category", "Sequence"]

def load_sheet_rows(ws, expected_columns):
    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    header_row = [h.strip() if isinstance(h, str) else h for h in header_row]
    col_index = {name: idx for idx, name in enumerate(header_row) if name in expected_columns}

    rows = []
    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in values):
            continue
        record = {}
        for col_name in expected_columns:
            idx = col_index.get(col_name)
            val = values[idx] if idx is not None and idx < len(values) else None
            if isinstance(val, str):
                val = val.strip()
            record[col_name] = val
        rows.append(record)
    return rows


def load_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    hymns = load_sheet_rows(wb[HYMNS_SHEET], HYMNS_COLUMNS)
    sections = load_sheet_rows(wb[SECTIONS_SHEET], SECTIONS_COLUMNS)
    contributors = load_sheet_rows(wb[CONTRIBUTORS_SHEET], CONTRIBUTORS_COLUMNS)
    categories = load_sheet_rows(wb[CATEGORIES_SHEET], CATEGORIES_COLUMNS)
    return hymns, sections, contributors, categories

def load_title_page(path):
    """Reads the 'Title Page' sheet (Key / Value columns) into a dict,
    e.g. {"title": "...", "subtitle": "..."}. Kept separate from
    load_workbook() so that function's return signature is unchanged."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[TITLE_PAGE_SHEET]
    info = {}
    for key, value in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        if key:
            info[str(key).strip().lower()] = value.strip() if isinstance(value, str) else value
    return info


def sort_key(title):
    """Strict alphabetical, case-insensitive, articles included; punctuation
    ignored (title 'Man of Sorrows!' and 'Man of Sorrows,' sort identically)."""
    t = (title or "").strip().lower()
    t = re.sub(r"[^\w\s]", " ", t)   # punctuation -> space, not deleted
    t = re.sub(r"\s+", " ", t).strip()
    return t




def normalize_lyric_text(text):
    """Strip stray leading/trailing whitespace on each authored line.
    ROOT-CAUSE FIX (confirmed via testing): a leading space after a
    line-internal newline in the workbook's Text cells (e.g.
    "line one\\n line two") was throwing off Typst's line alignment. The
    fix belongs here, at the data layer — not as a Typst-side .replace()
    workaround — so it's applied once, consistently, regardless of which
    presentation engine reads the data."""
    if not text:
        return text
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.strip() for line in text.split("\n"))


def group_contributors(contributors):
    """{hymn_id: [row, ...]}, each hymn's rows sorted by Sequence."""
    by_hymn = defaultdict(list)
    for c in contributors:
        by_hymn[c["Hymn ID"]].append(c)
    for hid in by_hymn:
        by_hymn[hid].sort(key=lambda c: (c["Sequence"] if isinstance(c["Sequence"], (int, float)) else 0))
    return by_hymn


def _role(row):
    return (row["Role"] or "").strip().lower()


def _rows_with_roles(rows, roles):
    return [r for r in rows if _role(r) in roles]


def _rows_for_tune_slot(rows, slot):
    return [r for r in _rows_with_roles(rows, TUNE_ROLES) if r.get("Tune Slot") in (slot, str(slot))]


# Each category's person-role/source-role pair, keyed by the role names used
# in Contributors. Used to split a category's rows into the "by" list
# (person roles) and the "from" list (source roles).
PERSON_ROLES = {"author", "translator", "composer"}
SOURCE_ROLES = {"word source", "translation source", "tune source"}


def join_credits(rows):
    """Oxford-comma name list from Contributors rows, each already in
    display order. A row's Note is appended in parens after that name only,
    e.g. 'A (Stanza 1), B (Stanza 2), and C'. Returns None for an empty list."""
    parts = []
    for r in rows:
        name = r["Person Name"]
        note = r.get("Note")
        parts.append(f"{name} ({note})" if note else name)

    if not parts:
        return None
    if len(parts) == 1:
        return parts[0]
    if len(parts) == 2:
        return f"{parts[0]} and {parts[1]}"
    return ", ".join(parts[:-1]) + f", and {parts[-1]}"


def by_from_suffix(rows):
    """Splits rows into person-role ('by') and source-role ('from') credits
    and returns the combined tail, e.g. ' by A and B & from C' — with a
    leading space, ready to append after a label or quoted tune name. Empty
    string if rows is empty. Established convention (2026-09-24): internal
    lists use Oxford-comma 'and'; the by/from halves are joined with ' & '
    only when both are present."""
    by_rows = [r for r in rows if _role(r) in PERSON_ROLES]
    from_rows = [r for r in rows if _role(r) in SOURCE_ROLES]
    by_join = join_credits(by_rows)
    from_join = join_credits(from_rows)

    if by_join and from_join:
        return f" by {by_join} & from {from_join}"
    if by_join:
        return f" by {by_join}"
    if from_join:
        return f" from {from_join}"
    return ""


def build_attribution(hymn, contrib_rows):
    """Returns (words_line, translator_line, tune_lines, alt_tunes_line, flags).
    tune_lines is a list (0 or 1 entries for this pilot's single-primary-tune
    model); alt_tunes_line is one consolidated string for everything else."""
    flags = []

    words_rows = _rows_with_roles(contrib_rows, WORDS_ROLES)
    translation_rows = _rows_with_roles(contrib_rows, TRANSLATION_ROLES)

    main_slot_raw = hymn.get("Main Tune")
    main_slot = int(main_slot_raw) if main_slot_raw not in (None, "") else 1

    tune_names = {1: hymn.get("Tune 1"), 2: hymn.get("Tune 2"), 3: hymn.get("Tune 3")}
    main_tune_name = tune_names.get(main_slot)
    main_composer_rows = _rows_for_tune_slot(contrib_rows, main_slot)

    if not words_rows:
        flags.append("No Author/Word Source in Contributors — 'Words by/from' line omitted, not invented")

    defaulted_to_title = not main_tune_name
    if defaulted_to_title:
        main_tune_name = hymn["Title"]
        flags.append(f"Main Tune (slot {main_slot}) name defaulted to hymn title (no separate tune name given)")

    # "Words and Tune by/from X" collapse: the Words-category credits and the
    # Main-Tune credits are the SAME SET of names (any size — one name or
    # several), all in the same "kind" (all person-role, or all source-role),
    # and none of them carry a Note. A person-credit matching a source-credit
    # by text alone is treated as a coincidence, not a collapse case, since
    # "by" vs "from" would otherwise be ambiguous for the single combined
    # line — so the two sides must also share the same kind.
    def _uniform_kind(rows):
        """True if every row is a person role, False if every row is a
        source role, None if rows is empty or mixes both kinds."""
        if not rows:
            return None
        kinds = {_role(r) in PERSON_ROLES for r in rows}
        return kinds.pop() if len(kinds) == 1 else None

    def _name_set(rows):
        return sorted((r["Person Name"] or "").strip().lower() for r in rows)

    w_kind = _uniform_kind(words_rows)
    c_kind = _uniform_kind(main_composer_rows)
    collapse = bool(
        words_rows and main_composer_rows
        and w_kind is not None and w_kind == c_kind
        and not any(r.get("Note") for r in words_rows + main_composer_rows)
        and _name_set(words_rows) == _name_set(main_composer_rows)
    )

    if collapse:
        connector = "by" if w_kind else "from"
        words_line = f"Words and Tune {connector} {join_credits(words_rows)}"
    elif words_rows:
        words_line = "Words" + by_from_suffix(words_rows)
    else:
        words_line = None

    translator_line = ("Translation" + by_from_suffix(translation_rows)) if translation_rows else None

    # Rule (2026-09-24): when a tune's name matches the hymn's own title
    # (case-insensitive — the workbook has titles in ALL CAPS in places, so
    # an exact-case comparison silently missed real matches), drop the
    # quoted name — "Tune by X" rather than 'Tune: "Title" by X'. Applies
    # to the Main Tune line only (not Alt. tune — see the alt-tune loop
    # below for why). This also covers the defaulted-to-title case above,
    # since main_tune_name there IS hymn Title by construction.
    tune_lines = []
    if not collapse:
        suffix = by_from_suffix(main_composer_rows)
        title_match = main_tune_name.strip().casefold() == (hymn.get("Title") or "").strip().casefold()
        tune_lines.append(("Tune" if title_match else f'Tune: "{main_tune_name}"') + suffix)
        if not suffix:
            flags.append(f"Main Tune '{main_tune_name}' has no Composer/Tune Source credited")

    # Alt. tune entries always keep their quoted name, even if it happened to
    # equal the hymn title — an alt tune sharing the hymn's own title would
    # be a real oddity, and dropping the name from just one part of a
    # semicolon-joined multi-alt line would be ambiguous about which entry
    # it belonged to.
    alt_parts = []
    for slot in (1, 2, 3):
        if slot == main_slot:
            continue
        name = tune_names.get(slot)
        if not name:
            continue
        suffix = by_from_suffix(_rows_for_tune_slot(contrib_rows, slot))
        alt_parts.append(f'"{name}"' + suffix)
        if not suffix:
            flags.append(f"Tune {slot} '{name}' has no Composer/Tune Source credited")

    alt_tunes_line = "Alt. tune: " + ";\n".join(alt_parts) if alt_parts else None

    return words_line, translator_line, tune_lines, alt_tunes_line, flags


def build_hymn_data(hymns, sections, contributors):
    by_hymn_sections = defaultdict(list)
    for s in sections:
        by_hymn_sections[s["Hymn ID"]].append(s)
    for hid in by_hymn_sections:
        by_hymn_sections[hid].sort(key=lambda s: s["Sequence"])

    by_hymn_contrib = group_contributors(contributors)

    active = [h for h in hymns if (h["Status"] or "").lower() == "active"]
    active.sort(key=lambda h: (sort_key(h["Title"]), h["ID"]))

    results = []
    for n, h in enumerate(active, start=1):
        contrib_rows = by_hymn_contrib.get(h["ID"], [])
        words_line, translator_line, tune_lines, alt_tunes_line, flags = build_attribution(h, contrib_rows)

        sections_out = []
        stanza_no = 0
        for s in by_hymn_sections.get(h["ID"], []):
            type_lc = (s["Type"] or "").strip().lower()
            entry = {"type": type_lc, "label": s["Label"] or None, "text": normalize_lyric_text(s["Text"])}
            if type_lc == "stanza":
                stanza_no += 1
                entry["stanza_no"] = stanza_no
            sections_out.append(entry)

        results.append({
            "printed_no": n,
            "id": h["ID"],
            "title": h["Title"],
            "category": h["Category"],
            "words_line": words_line,
            "translator_line": translator_line,
            "tune_lines": tune_lines,
            "alt_tunes_line": alt_tunes_line,
            "flags": flags,
            "sections": sections_out,
        })
    return results


INDEX3_ROLES = WORDS_ROLES + TRANSLATION_ROLES
INDEX4_ROLES = TUNE_ROLES


def build_indexes(hymns, hymn_data, contributors, categories):
    """hymns: raw Hymns Metadata rows (needed for Tune 1-3 / Main Tune).
    hymn_data: the printed-numbered active hymn dicts from build_hymn_data.
    contributors: raw Contributors rows.

    Returns (indexes_dict, warnings).

    DECIDED (2026-09-24, superseding the earlier last-name-sort plan): that
    plan was scrapped as too manual/volunteer-dependent. Index 3 and Index 4
    both sort exactly like the Title index — case/punctuation-insensitive,
    leading articles NOT ignored (sort_key(), no special-casing). No Persons
    lookup sheet is needed.

    Index 3 (combined Author/Source/Translator/Translation-Source): name ->
    sorted hymn numbers only. No role tag, no Note shown — a person or
    source appearing under multiple roles/hymns collapses to one entry with
    all their hymn numbers merged.

    Index 4 (Tune): every populated tune slot (main or alt) indexed under
    its own name -> {composer credit, hymn numbers}. Composer credit is a
    plain Oxford-comma name list (not the main-body "by/from" wording, and
    with Notes dropped — index entries stay compact), e.g.
    "Aberystwyth" - Joseph Parry - 229.

    Grouping key is (tune name, base-composer identity), NOT tune name
    alone (decided 2026-09-24, after two real collisions surfaced):
      - "base-composer identity" = the composer/tune-source row(s) for that
        slot explicitly noted "Stanza"/"Stanzas", when any are — NOT "every
        row except ones noted Chorus", because real data only tags the
        boundary rows of a role-group (e.g. Aurelia's actual Contributors
        rows are Wesley "(Stanzas)", then Getty/Getty/Cash with NO note at
        all, then de Barra "(Chorus)" — the middle rows aren't individually
        tagged, so an exclude-Chorus-only filter still let them leak into
        the identity and wrongly split the entry). Anchoring to an explicit
        Stanza(s) tag when present is robust to that partial tagging.
        A chorus added to a tune in some arrangements doesn't redefine
        whose tune it is, so "Aurelia" stays ONE index entry under Wesley
        regardless of which hymns also credit a chorus.
      - If NO row for that slot is Stanza-noted (the ordinary case — no
        stanza/chorus split at all), the full row set is used as-is, so a
        plain multi-composer tune (no notes) is unaffected by this rule.
      - Two hymns sharing a tune NAME but with a genuinely different
        base-composer identity are two DIFFERENT tunes that happen to
        share a name (e.g. "Depth Of Mercy") — they get separate index
        entries rather than being forced into one. A warning still
        surfaces this, since it usually deserves a human glance to confirm
        it's a real coincidence and not a typo.
      - Caveat: this relies on the base/stanza composer's row actually
        carrying the "(Stanzas)" note when a chorus is also credited on
        the same slot. If a future hymn tags only the chorus composers and
        leaves the base composer unnoted, this rule can't isolate it — tag
        the base composer's row "(Stanzas)" too (which is accurate anyway
        under the established Note convention) to fix it.
    """
    warnings = []
    by_hymn_contrib = group_contributors(contributors)
    hymns_by_id = {h["ID"]: h for h in hymns if h["ID"]}

    category_order = {c["Category"]: c["Sequence"] for c in categories if c.get("Category")}

    category_index = defaultdict(list)
    title_index = []

    person_source_nos = defaultdict(set)   # folded name -> {hymn_no, ...}
    person_source_display = {}             # folded name -> display text (first seen)

    tune_nos = defaultdict(set)             # (folded name, identity key) -> {hymn_no, ...}
    tune_display = {}                       # (folded name, identity key) -> [display name, credit]
    tune_identities_seen = defaultdict(set)  # folded name -> {identity key, ...} (for the collision warning)

    def _stanza_rows(rows):
        """Rows explicitly noted Stanza/Stanzas for this slot's credit —
        see build_indexes' docstring for why this anchors identity instead
        of trying to exclude Chorus-noted rows."""
        return [r for r in rows if r.get("Note") and re.search(r"(?i)\bstanzas?\b", r["Note"])]

    for h in hymn_data:
        entry = {"no": h["printed_no"], "title": h["title"]}
        category_index[h["category"] or "(uncategorized)"].append(entry)
        title_index.append(entry)

        contrib_rows = by_hymn_contrib.get(h["id"], [])

        for row in contrib_rows:
            if _role(row) not in INDEX3_ROLES:
                continue
            name = (row["Person Name"] or "").strip()
            if not name:
                continue
            key = name.casefold()
            person_source_nos[key].add(h["printed_no"])
            person_source_display.setdefault(key, name)

        raw = hymns_by_id.get(h["id"], {})
        main_slot_raw = raw.get("Main Tune")
        main_slot = int(main_slot_raw) if main_slot_raw not in (None, "") else 1

        for slot in (1, 2, 3):
            name = raw.get(f"Tune {slot}")
            if not name:
                if slot == main_slot:
                    name = h["title"]  # same default-to-title rule as attribution
                else:
                    continue
            name_key = name.casefold()

            credit_rows = [
                r for r in contrib_rows
                if _role(r) in INDEX4_ROLES and r.get("Tune Slot") in (slot, str(slot))
            ]
            base_rows = _stanza_rows(credit_rows)
            if not base_rows:
                base_rows = credit_rows  # no stanza/chorus split on this slot — use the full credit as-is

            identity = tuple(sorted((r["Person Name"] or "").strip().casefold() for r in base_rows))
            credit = join_credits([{**r, "Note": None} for r in base_rows]) or ""  # notes dropped in the index

            key = (name_key, identity)
            tune_nos[key].add(h["printed_no"])
            tune_display.setdefault(key, [name, credit])

            prior_identities = tune_identities_seen[name_key]
            if prior_identities and identity not in prior_identities:
                warnings.append(
                    f"Tune name '{name}' maps to more than one composer identity across hymns "
                    f"(now indexed as separate entries) — '{tune_display[key][1]}' vs credit(s) "
                    f"already seen for this name; confirm this is a genuine same-name coincidence, "
                    f"not a data-entry error (first seen with this credit at hymn #{h['printed_no']})"
                )
            prior_identities.add(identity)

    title_index.sort(key=lambda x: sort_key(x["title"]))

    person_source_index = {
        person_source_display[key]: sorted(nos)
        for key, nos in person_source_nos.items()
    }
    person_source_index = dict(
        sorted(person_source_index.items(), key=lambda kv: sort_key(kv[0]))
    )

    tune_index = [
        {"name": tune_display[key][0], "composer": tune_display[key][1], "hymn_nos": sorted(nos)}
        for key, nos in tune_nos.items()
    ]
    # Sort by tune name, then by composer as a tiebreaker so same-name
    # entries (a genuine collision, e.g. "Depth Of Mercy") stay adjacent
    # and print in a stable, deterministic order.
    tune_index.sort(key=lambda e: (sort_key(e["name"]), sort_key(e["composer"] or "")))

    category_out = dict(sorted(
        (
            (c, sorted(v, key=lambda x: sort_key(x["title"])))
            for c, v in category_index.items()
        ),
        key=lambda kv: category_order.get(kv[0], float("inf"))
    ))

    return {
        "category": category_out,
        "title": title_index,
        "person_source": person_source_index,
        "tune": tune_index,
    }, warnings


def print_summary(hymn_data):
    print("=" * 70)
    print(f"{len(hymn_data)} ACTIVE hymn(s) — printed order")
    print("=" * 70)
    for h in hymn_data:
        print(f"\n#{h['printed_no']:>3}  {h['title']}  [{h['id']}]  ({h['category']})")
        if h["words_line"]:
            print(f"      {h['words_line']}")
        if h["translator_line"]:
            print(f"      {h['translator_line']}")
        for line in h["tune_lines"]:
            print(f"      {line}")
        if h["alt_tunes_line"]:
            print(f"      {h['alt_tunes_line']}")
        types = ", ".join(s["type"] for s in h["sections"])
        print(f"      sections: {len(h['sections'])} ({types})")
        for f in h["flags"]:
            print(f"      \u26a0 {f}")

    print()
    print("=" * 70)
    print("Category counts")
    print("=" * 70)
    cats = defaultdict(int)
    for h in hymn_data:
        cats[h["category"] or "(uncategorized)"] += 1
    for c, n in sorted(cats.items()):
        print(f"  {c}: {n}")


def print_index_summary(indexes, warnings):
    print()
    print("=" * 70)
    print("Index summary")
    print("=" * 70)
    print(f"Category index : {len(indexes['category'])} categor{'y' if len(indexes['category']) == 1 else 'ies'}")
    print(f"Title index    : {len(indexes['title'])} entries")
    print(f"Author/Source/Translator index : {len(indexes['person_source'])} distinct name(s)")
    print(f"Tune index     : {len(indexes['tune'])} entries (a name may appear more than once if it's a genuine collision)")
    if warnings:
        print()
        print(f"Index warnings ({len(warnings)}):")
        for w in warnings:
            print(f"  ! {w}")


def main():
    if len(sys.argv) != 2:
        print("Usage: python3 transform.py <path-to-HYMNAL_EDITION_2.xlsx>")
        sys.exit(1)
    hymns, sections, contributors, categories = load_workbook(sys.argv[1])
    hymn_data = build_hymn_data(hymns, sections, contributors)
    print_summary(hymn_data)
    indexes, warnings = build_indexes(hymns, hymn_data, contributors, categories)
    print_index_summary(indexes, warnings)


if __name__ == "__main__":
    main()
